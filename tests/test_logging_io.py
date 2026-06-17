from datetime import datetime

from openpyxl import load_workbook

from sidm_monitor.logging_io import append_text_log, append_xlsx_log, log_path
from sidm_monitor.models import SessionRow


def test_log_path_uses_pb_date_and_server_name(tmp_path) -> None:
    path = log_path(tmp_path, "pfsvr0", datetime(2026, 6, 16, 9, 0, 0), ".txt")

    assert path == tmp_path / "2026.06.16_pfsvr0.txt"


def test_append_text_log_writes_start_marker_for_new_file(tmp_path) -> None:
    timestamp = datetime(2026, 6, 16, 9, 0, 0)

    path = append_text_log(tmp_path, "pfsvr0", "body\r\n", timestamp)

    assert path.read_text(encoding="utf-8").splitlines() == [
        "2026.06.16 09:00:00开始记录pfsvr0",
        "body",
    ]


def test_append_xlsx_log_creates_headers_and_rows(tmp_path) -> None:
    timestamp = datetime(2026, 6, 16, 9, 0, 0)
    rows = [
        SessionRow(
            tree="*62",
            flag="◆",
            session_id=62,
            blocking_session_id=0,
            host_name="server",
            ip_address="127.0.0.1",
            program_name="app",
            database_name="master",
            session_status="running",
            cpu_seconds=1,
            wait_seconds=12,
            wait_type="LCK_M_X",
            wait_resource="OBJECT:1",
            sql_text="select 1\r\n",
            kill_result="",
            login_name="sa",
        )
    ]

    path = append_xlsx_log(tmp_path, "pfsvr0", rows, timestamp)

    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet.cell(1, 1).value == "时间"
    assert sheet.cell(2, 1).value == "09:00:00"
    assert sheet.cell(2, 4).value == 62
    assert sheet.cell(2, 15).value == "select 1"
    assert sheet.cell(2, 17).value == "sa"
