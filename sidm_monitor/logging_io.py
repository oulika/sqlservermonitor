from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from sidm_monitor.models import SessionRow


XLSX_HEADERS = [
    "时间",
    "锁树",
    "锁标志",
    "会话",
    "锁会话",
    "主机名",
    "IP地址",
    "应用名",
    "数据库",
    "会话状态",
    "CPU时间",
    "等待时间",
    "等待类型",
    "等待的资源",
    "SQL语句",
    "中断",
    "登录名",
]


def log_path(log_dir: Path, server_name: str, timestamp: datetime, suffix: str) -> Path:
    return Path(log_dir) / f"{timestamp:%Y.%m.%d}_{server_name}{suffix}"


def append_text_log(
    log_dir: Path,
    server_name: str,
    text: str,
    timestamp: datetime,
) -> Path:
    path = log_path(log_dir, server_name, timestamp, ".txt")
    path.parent.mkdir(parents=True, exist_ok=True)
    is_new_file = not path.exists()

    with path.open("a", encoding="utf-8", newline="") as file:
        if is_new_file and "开始记录" not in text:
            file.write(f"{timestamp:%Y.%m.%d %H:%M:%S}开始记录{server_name}\r\n")
        file.write(text)

    return path


def append_xlsx_log(
    log_dir: Path,
    server_name: str,
    rows: Sequence[SessionRow],
    timestamp: datetime,
) -> Path:
    path = log_path(log_dir, server_name, timestamp, ".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(XLSX_HEADERS)

    for row in rows:
        if not row.tree:
            continue
        sheet.append(
            [
                f"{timestamp:%H:%M:%S}",
                row.tree,
                row.flag,
                row.session_id,
                row.blocking_session_id,
                row.host_name,
                row.ip_address,
                row.program_name,
                row.database_name,
                row.session_status,
                row.cpu_seconds,
                row.wait_seconds,
                row.wait_type,
                row.wait_resource,
                _one_line(row.sql_text),
                row.kill_result,
                row.login_name,
            ]
        )

    workbook.save(path)
    return path


def _one_line(text: str) -> str:
    return text.replace("\r", "").replace("\n", "")
